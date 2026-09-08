# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx==1.2.0"]
# ///
"""Build the synchronized second-edition Word, PDF, HTML, and EPUB textbook."""

from __future__ import annotations

import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
TEXTBOOK = ROOT / "textbook"
PUBLICATION = TEXTBOOK / "publishing"
EXPORTS = PUBLICATION / "exports"
FILTERS = PUBLICATION / "filters"
METADATA = PUBLICATION / "book_metadata.yaml"
REFERENCE_DOCX = PUBLICATION / "reference_2e.docx"
VERSION = "2.0.0-draft"
BASENAME = "operating_cloud_databases_second_edition_draft"

DOCX_PATH = EXPORTS / f"{BASENAME}.docx"
PDF_PATH = EXPORTS / f"{BASENAME}.pdf"
HTML_PATH = EXPORTS / f"{BASENAME}.html"
EPUB_PATH = EXPORTS / f"{BASENAME}.epub"

INK = "18222D"
MUTED = "526273"
TEAL = "0F766E"
TEAL_DARK = "0B4F4A"
BLUE = "315B7D"
PALE_TEAL = "DCEFEB"
CODE_FILL = "F2F7F6"
ORANGE = "B45309"

CONTENT_WIDTH_DXA = 9792  # 6.8 inches after 0.85-inch side margins.
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_SIDE_DXA = 120


def run(command: list[str], *, cwd: Path = ROOT, capture: bool = False) -> str:
    result = subprocess.run(
        command,
        cwd=cwd,
        check=True,
        text=capture,
        stdout=subprocess.PIPE if capture else None,
        stderr=subprocess.PIPE if capture else None,
    )
    return result.stdout if capture else ""


def require_tool(name: str, explicit: str | None = None) -> str:
    candidate = explicit or shutil.which(name)
    if not candidate or not Path(candidate).exists():
        raise RuntimeError(f"Required publication tool is missing: {name}")
    return candidate


def source_order() -> list[Path]:
    modules = {i: next(TEXTBOOK.glob(f"module_{i:02d}_*.md")) for i in range(1, 16)}
    return [
        PUBLICATION / "front_matter.md",
        PUBLICATION / "parts/part_1.md",
        modules[1], modules[2], modules[3],
        PUBLICATION / "parts/part_2.md",
        modules[4], modules[5], modules[6], modules[7], modules[8],
        PUBLICATION / "parts/part_3.md",
        modules[9], modules[10], modules[11], modules[12],
        PUBLICATION / "parts/part_4.md",
        modules[13], modules[14],
        PUBLICATION / "parts/part_5.md",
        modules[15],
        PUBLICATION / "appendix_a_notation_glossary.md",
        PUBLICATION / "appendix_b_technical_templates.md",
        PUBLICATION / "back_matter.md",
    ]


def render_visual_sources() -> None:
    dot = require_tool("dot")
    magick = require_tool("magick")

    for source in sorted((TEXTBOOK / "figures").glob("*.dot")):
        target = source.with_suffix(".png")
        if not target.exists() or source.stat().st_mtime > target.stat().st_mtime:
            run([dot, "-Tpng", "-Gdpi=220", str(source), "-o", str(target)])

    for source in sorted((TEXTBOOK / "figures").glob("*.svg")):
        target = source.with_suffix(".png")
        if not target.exists() or source.stat().st_mtime > target.stat().st_mtime:
            run([magick, "-background", "none", str(source), str(target)])

    cover_svg_path = PUBLICATION / "cover_2e.svg"
    cover_png_path = PUBLICATION / "cover_2e.png"
    original_cover = cover_svg_path.read_text(encoding="utf-8")
    cover_svg = original_cover.replace(
        "Version 1.0.0-rc.1  |  CC BY-NC-SA 4.0",
        "Second edition draft  |  CC BY-NC-SA 4.0",
    ).replace(
        "A teal and cream cover with connected relational tables, documents, and cloud database symbols.",
        "A teal and cream second-edition cover connecting relational, document, and managed cloud database ideas.",
    ).replace(
        "service + evidence",
        "service + operations",
    ).replace(
        'font-size="46" font-weight="650">A LAB-FIRST COURSE BUILT AROUND EVIDENCE',
        'font-size="42" font-weight="650">A LAB-FIRST COURSE IN DATABASE REASONING AND OPERATIONS',
    )
    cover_svg_path.write_text(cover_svg, encoding="utf-8")
    run([magick, "-background", "none", str(cover_svg_path), str(cover_png_path)])


def color(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_style_font(style, name: str, size: float, hex_color: str = INK, *, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color(hex_color)
    if bold is not None:
        style.font.bold = bold

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attribute in list(rfonts.attrib):
        if "theme" in attribute.lower():
            del rfonts.attrib[attribute]
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), name)


def ensure_style(document: Document, name: str, style_type: WD_STYLE_TYPE, base: str | None = None):
    try:
        style = document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, style_type)
    if base:
        style.base_style = document.styles[base]
    return style


def add_paragraph_bottom_border(paragraph, hex_color: str, size: int = 8) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), hex_color)


def append_page_field(paragraph) -> None:
    paragraph.add_run("SECOND EDITION DRAFT  |  PAGE ")
    run_element = paragraph.add_run()._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_element.extend([begin, instruction, separate, display, end])


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.text = "OPERATING CLOUD DATABASES  |  SECOND EDITION DRAFT"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    for item in p.runs:
        item.font.name = "Linux Biolinum G"
        item.font.size = Pt(8)
        item.font.color.rgb = color(MUTED)
    add_paragraph_bottom_border(p, PALE_TEAL, 6)

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""

    footer = section.footer
    fp = footer.paragraphs[0]
    clear_paragraph(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(2)
    append_page_field(fp)
    for item in fp.runs:
        item.font.name = "Linux Biolinum G"
        item.font.size = Pt(8)
        item.font.color.rgb = color(MUTED)

    section.first_page_footer.paragraphs[0].text = ""


def build_reference_docx(pandoc: str) -> None:
    binary = subprocess.run(
        [pandoc, "--print-default-data-file", "reference.docx"],
        cwd=ROOT,
        check=True,
        stdout=subprocess.PIPE,
    ).stdout
    REFERENCE_DOCX.write_bytes(binary)

    document = Document(REFERENCE_DOCX)
    styles = document.styles

    normal = styles["Normal"]
    set_style_font(normal, "Linux Libertine G", 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.widow_control = True

    for body_name in ("Body Text", "Block Text", "First Paragraph"):
        if body_name in styles:
            set_style_font(styles[body_name], "Linux Libertine G", 11)
            styles[body_name].paragraph_format.space_after = Pt(6)
            styles[body_name].paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    set_style_font(title, "Linux Biolinum G", 30, TEAL_DARK, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, "Linux Biolinum G", 15, BLUE)
    subtitle.paragraph_format.space_after = Pt(16)

    heading_tokens = {
        "Heading 1": (19, TEAL_DARK, 18, 10),
        "Heading 2": (14, TEAL, 14, 7),
        "Heading 3": (11.5, BLUE, 10, 5),
        "Heading 4": (10.5, MUTED, 8, 4),
    }
    for name, (size, tone, before, after) in heading_tokens.items():
        style = styles[name]
        set_style_font(style, "Linux Biolinum G", size, tone, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True

    part = ensure_style(document, "Part Title", WD_STYLE_TYPE.PARAGRAPH, "Heading 1")
    set_style_font(part, "Linux Biolinum G", 24, "FFFFFF", bold=True)
    part.paragraph_format.space_before = Pt(0)
    part.paragraph_format.space_after = Pt(18)
    part.paragraph_format.page_break_before = True
    part.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    set_style_font(caption, "Linux Biolinum G", 8.5, MUTED)
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    listing = ensure_style(document, "Listing Caption", WD_STYLE_TYPE.PARAGRAPH, "Caption")
    set_style_font(listing, "Linux Biolinum G", 8.5, MUTED, bold=True)
    listing.font.italic = False
    listing.paragraph_format.space_before = Pt(8)
    listing.paragraph_format.space_after = Pt(2)
    listing.paragraph_format.keep_with_next = True

    source = ensure_style(document, "Source Code", WD_STYLE_TYPE.PARAGRAPH, "Normal")
    set_style_font(source, "Liberation Mono", 9, INK)
    source.paragraph_format.left_indent = Inches(0.14)
    source.paragraph_format.right_indent = Inches(0.06)
    source.paragraph_format.space_before = Pt(2)
    source.paragraph_format.space_after = Pt(7)
    source.paragraph_format.line_spacing = 1.05
    source.paragraph_format.keep_together = True

    verbatim = ensure_style(document, "Verbatim Char", WD_STYLE_TYPE.CHARACTER)
    set_style_font(verbatim, "Liberation Mono", 8.5, INK)

    for list_name in ("List Bullet", "List Number", "List Continue"):
        if list_name in styles:
            style = styles[list_name]
            set_style_font(style, "Linux Libertine G", 11, INK)
            style.paragraph_format.left_indent = Inches(0.375)
            style.paragraph_format.first_line_indent = Inches(-0.188)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.line_spacing = 1.18

    for section in document.sections:
        configure_section(section)

    document.core_properties.title = "Operating Cloud Databases"
    document.core_properties.subject = "Second edition draft"
    document.core_properties.author = "Atilio Barreda"
    document.core_properties.keywords = "database administration, SQL, PostgreSQL, MongoDB, cloud databases, OER"
    document.core_properties.comments = "Unpublished second-edition draft; generated from editable source."
    document.save(REFERENCE_DOCX)


def add_shading(element, fill: str) -> None:
    paragraph_element = element._p if hasattr(element, "_p") else element
    properties = paragraph_element.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def add_left_border(paragraph, tone: str = TEAL) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), tone)


def replace_with_toc_field(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run_element = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = "Update this table of contents in Word if page numbers change."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_element.extend([begin, instruction, separate, text_node, end])
    paragraph._p.append(run_element)


def pandoc_identifier(title: str) -> str:
    value = re.sub(r"[`*_]", "", title).lower()
    value = re.sub(r"[^\w\s-]", "", value, flags=re.UNICODE)
    value = re.sub(r"[\s_]+", "-", value).strip("-")
    return value


def first_level_one_heading(path: Path) -> str:
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            heading = re.sub(r"\s+\{[^}]*\}\s*$", "", line[2:]).strip()
            return heading
    raise RuntimeError(f"No level-one heading found in {path}")


def contents_entries() -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    for path in source_order():
        title = first_level_one_heading(path)
        kind = "chapter"
        label = title
        if path.parent.name == "parts":
            kind = "part"
        elif path.name == "front_matter.md":
            kind = "front"
        elif path.name.startswith("appendix_"):
            kind = "appendix"
        elif path.name == "back_matter.md":
            kind = "back"
        elif path.parent == TEXTBOOK:
            number = int(path.name.split("_")[1])
            label = f"{number}. {title}"
        entries.append({
            "title": title,
            "label": label,
            "identifier": pandoc_identifier(title),
            "kind": kind,
        })
    return entries


def render_static_contents(page_map: dict[str, int]) -> str:
    lines = [
        "# Contents {.unnumbered}",
        "",
        "The page numbers below are generated from the same Word layout used for this PDF.",
        "",
        "| Section | Page |",
        "|---|---:|",
    ]
    for entry in contents_entries():
        page = str(page_map.get(entry["title"], "-"))
        link = f"[{entry['label']}](#{entry['identifier']})"
        if entry["kind"] == "part":
            link = f"**{link}**"
            page = f"**{page}**"
        lines.append(f"| {link} | {page} |")
    lines.extend(["", "::: {.pagebreak}", ":::", ""])
    return "\n".join(lines)


def extract_contents_pages(pdf_path: Path, pdftotext: str) -> dict[str, int]:
    text = run([pdftotext, "-layout", str(pdf_path), "-"], capture=True)
    pages = text.split("\f")
    entries = contents_entries()
    normalized_pages = [re.sub(r"\s+", " ", page).strip() for page in pages]

    # A contents page mentions many chapter titles; exclude it from title matching.
    content_page_indexes = []
    for index, page in enumerate(normalized_pages):
        title_matches = sum(entry["title"] in page for entry in entries)
        if title_matches < 5:
            content_page_indexes.append(index)

    result: dict[str, int] = {}
    for entry in entries:
        for index in content_page_indexes:
            if entry["title"] in normalized_pages[index]:
                result[entry["title"]] = index + 1
                break
    return result


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    values = {
        "top": CELL_TOP_BOTTOM_DXA,
        "bottom": CELL_TOP_BOTTOM_DXA,
        "start": CELL_SIDE_DXA,
        "end": CELL_SIDE_DXA,
    }
    for edge, amount in values.items():
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_row_repeat_and_keep(row, *, repeat: bool) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if repeat and tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def choose_column_widths(table) -> list[int]:
    column_count = len(table.columns)
    if [cell.text.strip() for cell in table.rows[0].cells] == ["Section", "Page"]:
        return [CONTENT_WIDTH_DXA - 800, 800]
    maxima: list[int] = []
    for column_index in range(column_count):
        lengths = []
        for row in table.rows:
            cell = row.cells[column_index]
            # python-docx cell.text omits equation runs; include them in sizing.
            content = " ".join(cell._tc.xpath(".//w:t/text() | .//m:t/text()"))
            math_length = sum(len(value) for value in cell._tc.xpath(".//m:t/text()"))
            lengths.append(len(re.sub(r"\s+", " ", content).strip()) + math_length // 3)
        maximum = max(lengths, default=10)
        maxima.append(max(10, min(70, maximum)))

    natural_widths = [max(780, maximum * 105 + 240) for maximum in maxima]
    if sum(natural_widths) <= CONTENT_WIDTH_DXA:
        return natural_widths

    minimum = 900 if column_count >= 4 else 1050
    if minimum * column_count >= CONTENT_WIDTH_DXA:
        minimum = CONTENT_WIDTH_DXA // column_count
    remaining = CONTENT_WIDTH_DXA - (minimum * column_count)
    weight_total = sum(maxima) or column_count
    widths = [minimum + int(remaining * weight / weight_total) for weight in maxima]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_table_geometry(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = choose_column_widths(table)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")

    grid = table._tbl.tblGrid
    for grid_col in list(grid):
        grid.remove(grid_col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row_index, row in enumerate(table.rows):
        set_row_repeat_and_keep(row, repeat=row_index == 0)
        for column_index, cell in enumerate(row.cells):
            cell.width = Twips(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[column_index]))
            tc_w.set(qn("w:type"), "dxa")
            if row_index == 0:
                shade_cell(cell, PALE_TEAL)
            elif row_index % 2 == 0:
                shade_cell(cell, "F5F7F7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for item in paragraph.runs:
                    is_code = item.style and item.style.name == "Verbatim Char"
                    item.font.name = "Liberation Mono" if is_code else ("Linux Biolinum G" if row_index == 0 else "Linux Libertine G")
                    item.font.size = Pt(9 if is_code else (9.5 if row_index == 0 else 10))
                    if row_index == 0:
                        item.bold = True
                        item.font.color.rgb = color("173652")


def polish_docx(path: Path) -> None:
    document = Document(path)
    math_properties = document.settings.element.find(qn("m:mathPr"))
    if math_properties is None:
        math_properties = OxmlElement("m:mathPr")
        document.settings.element.append(math_properties)
    math_font = math_properties.find(qn("m:mathFont"))
    if math_font is None:
        math_font = OxmlElement("m:mathFont")
        math_properties.insert(0, math_font)
    math_font.set(qn("m:val"), "DejaVu Math TeX Gyre")
    for section in document.sections:
        configure_section(section)

    update_fields = document.settings.element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        document.settings.element.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        value = paragraph.text.strip()
        if value == "[[TOC_FIELD]]":
            replace_with_toc_field(paragraph)
        elif value.startswith("Listing "):
            paragraph.style = document.styles["Listing Caption"]
        elif re.match(r"^Part [IVX]+:", value):
            paragraph.style = document.styles["Part Title"]
            add_shading(paragraph, TEAL_DARK)
            add_left_border(paragraph, ORANGE)

        if paragraph.style and paragraph.style.name == "Source Code":
            add_shading(paragraph, CODE_FILL)
            add_left_border(paragraph, TEAL)
            paragraph.paragraph_format.keep_together = True
        elif value.endswith(":") and len(value) <= 120:
            paragraph.paragraph_format.keep_with_next = True

        has_drawing = bool(paragraph._p.xpath(".//w:drawing"))
        if has_drawing:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(2)
            if index + 1 < len(paragraphs) and paragraphs[index + 1].style.name == "Caption":
                paragraph.paragraph_format.keep_with_next = True

    for table in document.tables:
        set_table_geometry(table)

    for shape in document.inline_shapes:
        doc_pr = shape._inline.docPr
        if not doc_pr.get("descr"):
            doc_pr.set("descr", "Textbook figure; see the adjacent caption for a full description.")

    document.core_properties.title = "Operating Cloud Databases"
    document.core_properties.subject = "PostgreSQL, MongoDB, and Reliable Data Systems - second edition draft"
    document.core_properties.author = "Atilio Barreda"
    document.core_properties.keywords = "database administration, relational algebra, SQL, PostgreSQL, Supabase, MongoDB, Atlas, OER"
    document.core_properties.comments = "Generated from the editable second-edition source; unpublished draft."
    document.save(path)


def pandoc_base(pandoc: str) -> list[str]:
    resource_path = os.pathsep.join([str(TEXTBOOK), str(PUBLICATION), str(ROOT)])
    return [
        pandoc,
        "--from=markdown+fenced_divs+tex_math_dollars+pipe_tables+implicit_figures+link_attributes+bracketed_spans+header_attributes",
        "--number-sections",
        "--toc-depth=2",
        f"--resource-path={resource_path}",
        f"--lua-filter={FILTERS / 'pagebreak.lua'}",
        f"--lua-filter={FILTERS / 'companion_links.lua'}",
        f"--lua-filter={FILTERS / 'listings.lua'}",
        f"--highlight-style={PUBLICATION / 'syntax.theme'}",
    ]


def build_docx(
    pandoc: str,
    temporary_directory: Path,
    page_map: dict[str, int],
) -> None:
    cover_markdown = temporary_directory / "docx_cover.md"
    cover_markdown.write_text(
        "![](textbook/publishing/cover_2e.png){width=5.45in}\n\n"
        "::: {.pagebreak}\n:::\n",
        encoding="utf-8",
    )
    contents_markdown = temporary_directory / "docx_contents.md"
    contents_markdown.write_text(render_static_contents(page_map), encoding="utf-8")
    command = pandoc_base(pandoc) + [
        f"--reference-doc={REFERENCE_DOCX}",
        "--output",
        str(DOCX_PATH),
        str(cover_markdown),
        str(contents_markdown),
        *[str(path) for path in source_order()],
    ]
    run(command)
    polish_docx(DOCX_PATH)


def build_html(pandoc: str) -> None:
    command = pandoc_base(pandoc) + [
        "--standalone",
        "--toc",
        "--section-divs",
        "--mathml",
        "--embed-resources",
        f"--metadata-file={METADATA}",
        f"--template={PUBLICATION / 'template.html'}",
        f"--css={PUBLICATION / 'operating_cloud_databases_2e.css'}",
        "--output",
        str(HTML_PATH),
        *[str(path) for path in source_order()],
    ]
    run(command)


def build_epub(pandoc: str) -> None:
    command = pandoc_base(pandoc) + [
        "--standalone",
        "--toc",
        "--mathml",
        f"--metadata-file={METADATA}",
        f"--css={PUBLICATION / 'operating_cloud_databases_2e.css'}",
        f"--epub-cover-image={PUBLICATION / 'cover_2e.png'}",
        "--output",
        str(EPUB_PATH),
        *[str(path) for path in source_order()],
    ]
    run(command)


def build_pdf(soffice: str, temporary_directory: Path, pass_number: int) -> None:
    profile = temporary_directory / f"libreoffice-profile-{pass_number}"
    profile.mkdir(parents=True, exist_ok=True)
    output_directory = temporary_directory / f"pdf-pass-{pass_number}"
    output_directory.mkdir(parents=True, exist_ok=True)
    run([
        soffice,
        f"-env:UserInstallation={profile.as_uri()}",
        "--headless",
        "--convert-to",
        'pdf:writer_pdf_Export:{"UseTaggedPDF":{"type":"boolean","value":"true"},"ExportBookmarks":{"type":"boolean","value":"true"}}',
        "--outdir",
        str(output_directory),
        str(DOCX_PATH),
    ])
    generated = output_directory / f"{DOCX_PATH.stem}.pdf"
    shutil.copy2(generated, PDF_PATH)


def write_manifests() -> None:
    outputs = [DOCX_PATH, PDF_PATH, HTML_PATH, EPUB_PATH]
    checksums = []
    for path in outputs:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        checksums.append(f"{digest}  {path.name}")
    (EXPORTS / "SHA256SUMS.txt").write_text("\n".join(checksums) + "\n", encoding="utf-8")

    sources = []
    for path in source_order():
        sources.append({
            "path": str(path.relative_to(ROOT)),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        })
    manifest = {
        "title": "Operating Cloud Databases",
        "version": VERSION,
        "status": "unpublished second-edition draft",
        "canonical_sources": sources,
        "generated_outputs": [path.name for path in outputs],
        "supporting_sources": [
            {"path": str(path.relative_to(ROOT)),
             "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
            for path in sorted({
                METADATA, Path(__file__).resolve(),
                PUBLICATION / "cover_2e.svg", PUBLICATION / "cover_2e.png",
                PUBLICATION / "operating_cloud_databases_2e.css",
                PUBLICATION / "template.html",
                PUBLICATION / "syntax.theme",
                *FILTERS.glob("*.lua"),
                *(path for path in (TEXTBOOK / "figures").rglob("*")
                  if path.suffix in {".png", ".svg", ".dot"}),
            }) if path.is_file()
        ],
    }
    (EXPORTS / "SOURCE_MANIFEST.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )


def structural_checks() -> None:
    expected = [DOCX_PATH, PDF_PATH, HTML_PATH, EPUB_PATH]
    for path in expected:
        if not path.exists() or path.stat().st_size < 10_000:
            raise RuntimeError(f"Publication output is missing or unexpectedly small: {path}")

    html_text = HTML_PATH.read_text(encoding="utf-8")
    checks = {
        "skip link": 'class="skip-link"',
        "main landmark": 'id="main-content"',
        "MathML": "<math",
        "numbered listing": "Listing 1.",
        "relational selection symbol": "σ",
        "cloud interface figure": "A redacted MongoDB Atlas project overview",
    }
    missing = [label for label, needle in checks.items() if needle not in html_text]
    if missing:
        raise RuntimeError("HTML structural checks failed: " + ", ".join(missing))

    chapter_count = len(re.findall(r'<h1[^>]*data-number="(?:[0-9]+)"', html_text))
    if chapter_count < 15:
        raise RuntimeError(f"Expected at least 15 numbered chapters; found {chapter_count}")

    print("Built and structurally checked:")
    for path in expected:
        print(f"  {path.relative_to(ROOT)} ({path.stat().st_size:,} bytes)")


def main() -> None:
    EXPORTS.mkdir(parents=True, exist_ok=True)
    pandoc = require_tool("pandoc")
    soffice = require_tool(
        "soffice",
        os.environ.get("SOFFICE"),
    )
    pdftotext = require_tool("pdftotext")

    render_visual_sources()
    build_reference_docx(pandoc)
    with tempfile.TemporaryDirectory(prefix="cst4714-publication-2e-") as temporary:
        temporary_directory = Path(temporary)
        page_map: dict[str, int] = {}
        for pass_number in range(1, 6):
            build_docx(pandoc, temporary_directory, page_map)
            build_pdf(soffice, temporary_directory, pass_number)
            observed_page_map = extract_contents_pages(PDF_PATH, pdftotext)
            missing = [
                entry["title"]
                for entry in contents_entries()
                if entry["title"] not in observed_page_map
            ]
            if missing:
                raise RuntimeError(
                    "Could not locate these top-level headings in the PDF: "
                    + ", ".join(missing)
                )
            if page_map and observed_page_map == page_map:
                break
            page_map = observed_page_map
        else:
            raise RuntimeError("The generated contents page numbers did not stabilize.")

        build_html(pandoc)
        build_epub(pandoc)
    write_manifests()
    structural_checks()


if __name__ == "__main__":
    main()
