#!/usr/bin/env python3
"""Build the Fall 2026 student syllabus DOCX from the canonical Markdown."""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "syllabus.md"
OUTPUT = ROOT / "CST4714_Fall_2026_Syllabus.docx"

BLUE = "2E5D95"
DARK_BLUE = "1F4E79"
LIGHT_BLUE = "EAF0F7"
LIGHT_GRAY = "F3F3F3"
LINE_GRAY = "C9CED6"
MID_GRAY = "666666"
BLACK = "1F1F1F"
WHITE = "FFFFFF"
BODY_FONT = "Georgia"
HEADING_FONT = "Arial"


def set_font(run, *, name=BODY_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=LINE_GRAY, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, placeholder, end])
    set_font(run, name=HEADING_FONT, size=8, color=MID_GRAY)


def set_paragraph_bottom_border(paragraph, color=BLUE, size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def style_callout(paragraph):
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    for run in paragraph.runs:
        set_font(run, size=10, italic=True, color=MID_GRAY)


def configure_styles(doc):
    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph"):
        style = styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style.font.size = Pt(10.2)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.06

    compact = styles["Compact"]
    compact.font.name = BODY_FONT
    compact._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    compact._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    compact.font.size = Pt(10)
    compact.font.color.rgb = RGBColor.from_string(BLACK)
    compact.paragraph_format.space_after = Pt(2)
    compact.paragraph_format.line_spacing = 1.04

    block = styles["Block Text"]
    block.font.name = BODY_FONT
    block.font.size = Pt(10)
    block.font.color.rgb = RGBColor.from_string(MID_GRAY)
    block.paragraph_format.line_spacing = 1.04

    for level, size, color, before, after in (
        (1, 18, BLUE, 12, 5),
        (2, 13.5, BLUE, 11, 4),
        (3, 11, DARK_BLUE, 8, 3),
    ):
        style = styles[f"Heading {level}"]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Hyperlink" in styles:
        hyperlink = styles["Hyperlink"]
        hyperlink.font.color.rgb = RGBColor.from_string(BLUE)
        hyperlink.font.underline = True

    for name in ("Verbatim Char", "Source Code"):
        if name in styles:
            styles[name].font.name = "Consolas"
            styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")


def configure_section(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.95)
        section.right_margin = Inches(0.95)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        section.different_first_page_header_footer = True

        header = section.header.paragraphs[0]
        header.paragraph_format.space_after = Pt(0)
        set_font(
            header.add_run("CST4714 Database Administration  |  Fall 2026"),
            name=HEADING_FONT,
            size=8,
            color=MID_GRAY,
        )

        for footer_part in (section.footer, section.first_page_footer):
            footer = footer_part.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer.paragraph_format.space_before = Pt(0)
            set_font(
                footer.add_run("Page "),
                name=HEADING_FONT,
                size=8,
                color=MID_GRAY,
            )
            add_page_field(footer, "PAGE")
            set_font(
                footer.add_run(" of "),
                name=HEADING_FONT,
                size=8,
                color=MID_GRAY,
            )
            add_page_field(footer, "NUMPAGES")


def style_paragraphs(doc):
    if len(doc.paragraphs) < 3:
        raise RuntimeError("Syllabus conversion produced too few paragraphs")

    title = doc.paragraphs[0]
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(1)
    for run in title.runs:
        set_font(run, name=HEADING_FONT, size=22, bold=True, color=BLUE)

    subtitle = doc.paragraphs[1]
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(9)
    set_paragraph_bottom_border(subtitle)
    for run in subtitle.runs:
        set_font(run, name=HEADING_FONT, size=12.5, bold=False, color=MID_GRAY)

    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            if paragraph.text == "Late Work and Technical Problems":
                paragraph.paragraph_format.page_break_before = True
        elif paragraph.style.name in {"Body Text", "First Paragraph", "Normal"}:
            paragraph.paragraph_format.widow_control = True
            if index and doc.paragraphs[index - 1].style.name == "Compact":
                paragraph.paragraph_format.space_before = Pt(5)

        if paragraph.style.name == "Block Text":
            style_callout(paragraph)


def style_tables(doc):
    geometries = (
        ([1900, 7460], 9.2),
        ([2500, 900, 5960], 8.9),
        ([1100, 3580, 1100, 3580], 9.2),
        ([800, 4900, 3660], 8.4),
        ([1200, 3480, 1200, 3480], 8.2),
    )
    if len(doc.tables) != len(geometries):
        raise RuntimeError(f"Expected 5 syllabus tables, found {len(doc.tables)}")

    for table_index, (table, (widths, font_size)) in enumerate(zip(doc.tables, geometries)):
        set_table_borders(table)
        set_table_geometry(table, widths)
        mark_header_row(table.rows[0])

        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell_index, cell in enumerate(row.cells):
                if row_index == 0:
                    shade_cell(cell, LIGHT_BLUE)
                elif table_index == 0 and cell_index == 0:
                    shade_cell(cell, LIGHT_GRAY)
                else:
                    shade_cell(cell, WHITE)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.02
                    if table_index == 1 and cell_index == 1:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if table_index == 2:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if table_index == 3 and cell_index == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if table_index == 4 and cell_index in {0, 2}:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        is_header = row_index == 0
                        is_label = table_index == 0 and cell_index == 0
                        is_week = table_index == 3 and cell_index == 0
                        is_date = table_index == 4 and cell_index in {0, 2}
                        set_font(
                            run,
                            name=HEADING_FONT if is_header else BODY_FONT,
                            size=font_size,
                            bold=(is_header or is_label or is_week or is_date),
                            color=DARK_BLUE if is_header else BLACK,
                        )


def build():
    with tempfile.TemporaryDirectory(prefix="cst4714_syllabus_") as temp_dir:
        raw_docx = Path(temp_dir) / "raw.docx"
        subprocess.run(
            [
                "/opt/homebrew/bin/pandoc",
                str(SOURCE),
                "--from=gfm",
                "--to=docx",
                f"--output={raw_docx}",
            ],
            check=True,
        )

        doc = Document(raw_docx)
        configure_styles(doc)
        configure_section(doc)
        style_paragraphs(doc)
        style_tables(doc)

        properties = doc.core_properties
        properties.title = "CST4714 Database Administration - Fall 2026 Syllabus"
        properties.subject = "Fall 2026 course syllabus"
        properties.author = "Professor Atilio Barreda"
        properties.keywords = "CST4714, database administration, PostgreSQL, MongoDB, Fall 2026, City Tech"
        properties.comments = ""

        doc.save(OUTPUT)
        print(OUTPUT)


if __name__ == "__main__":
    build()
